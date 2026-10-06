# -*- coding: utf-8 -*-
"""文件格式转换。依赖的开源项目:Pillow(图片)、FFmpeg(音视频,经 imageio-ffmpeg 内置)、
pypdf / pypdfium2(PDF)、python-docx(Word)、Qt 文本排版引擎(文字 → PDF)。"""
import os, subprocess
from .paths import unique_path, no_window_kwargs

IMG = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif', '.tif', '.tiff', '.ico', '.icns', '.ppm', '.tga'}
DOC = {'.txt', '.md', '.markdown', '.docx', '.pdf'}
VID = {'.mp4', '.mov', '.mkv', '.avi', '.webm', '.flv', '.wmv', '.m4v', '.mpg', '.mpeg', '.ts', '.3gp'}
AUD = {'.mp3', '.wav', '.flac', '.m4a', '.aac', '.ogg', '.opus', '.wma'}
ARC = {'.zip', '.7z', '.rar', '.tar', '.gz', '.tgz', '.bz2', '.xz'}

TARGETS = {
    'image': ['png', 'jpg', 'webp', 'ico', 'bmp', 'gif', 'tiff', 'icns', 'pdf'],
    'video': ['gif', 'mp4', 'webm', 'mov', 'mkv', 'avi', 'mp3'],
    'audio': ['mp3', 'wav', 'flac', 'm4a', 'ogg'],
}
DOC_TARGETS = {'.txt': ['pdf', 'docx', 'md'], '.md': ['pdf', 'docx', 'txt'], '.markdown': ['pdf', 'docx', 'txt'],
               '.docx': ['pdf', 'txt', 'md'], '.pdf': ['txt', 'png', 'jpg']}


def category(path):
    if os.path.isdir(path):
        return 'folder'
    e = os.path.splitext(path)[1].lower()
    for name, exts in (('image', IMG), ('doc', DOC), ('video', VID), ('audio', AUD), ('archive', ARC)):
        if e in exts:
            return name
    return 'other'


def targets_for(paths):
    """这批文件共同可用的目标格式。"""
    cats = {category(p) for p in paths}
    if len(cats) != 1:
        return []
    cat = cats.pop()
    if cat == 'doc':
        sets = [DOC_TARGETS.get(os.path.splitext(p)[1].lower(), []) for p in paths]
        return [t for t in sets[0] if all(t in s for s in sets)]
    out = list(TARGETS.get(cat, []))
    if cat == 'image' and all(p.lower().endswith('.gif') for p in paths):
        out.append('mp4')
    exts = {os.path.splitext(p)[1].lower().lstrip('.').replace('jpeg', 'jpg') for p in paths}
    if len(exts) == 1:
        out = [t for t in out if t not in exts]
    return out


def ffmpeg_exe():
    import imageio_ffmpeg
    return imageio_ffmpeg.get_ffmpeg_exe()


def _run_ffmpeg(args):
    p = subprocess.run([ffmpeg_exe(), '-y', '-hide_banner', '-loglevel', 'error'] + args,
                       capture_output=True, **no_window_kwargs())
    if p.returncode != 0:
        raise RuntimeError('FFmpeg 转换失败:' + p.stderr.decode('utf-8', 'ignore')[-300:])


def _out(src, fmt):
    return unique_path(os.path.splitext(src)[0] + '.' + fmt)


def convert_image(src, fmt):
    from PIL import Image
    if fmt == 'mp4':
        return convert_media(src, fmt)
    out = _out(src, fmt)
    im = Image.open(src)
    if fmt in ('jpg', 'bmp', 'pdf'):
        if im.mode in ('RGBA', 'LA', 'P'):
            im = im.convert('RGBA')
            bg = Image.new('RGB', im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert('RGB')
    if fmt == 'jpg':
        im.save(out, 'JPEG', quality=92, optimize=True)
    elif fmt == 'webp':
        im.save(out, 'WEBP', quality=90, method=6)
    elif fmt == 'ico':
        im = im.convert('RGBA')
        side = max(im.size)
        sq = Image.new('RGBA', (side, side), (0, 0, 0, 0))
        sq.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        sq.save(out, 'ICO', sizes=[(s, s) for s in (16, 24, 32, 48, 64, 128, 256) if s <= max(side, 16)])
    elif fmt == 'icns':
        im = im.convert('RGBA')
        side = max(im.size)
        sq = Image.new('RGBA', (side, side), (0, 0, 0, 0))
        sq.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        sq.resize((1024, 1024), Image.LANCZOS).save(out, 'ICNS')
    elif fmt == 'tiff':
        im.save(out, 'TIFF', compression='tiff_lzw')
    elif fmt == 'gif':
        im.save(out, 'GIF', save_all=getattr(im, 'is_animated', False))
    else:
        im.save(out, fmt.upper())
    return out


def convert_media(src, fmt):
    out = _out(src, fmt)
    i = ['-i', src]
    if fmt == 'gif':
        vf = ("fps=15,scale='min(800,iw)':-2:flags=lanczos,split[a][b];[a]palettegen=stats_mode=diff[p];"
              "[b][p]paletteuse=dither=sierra2_4a")
        _run_ffmpeg(i + ['-vf', vf, '-loop', '0', out])
    elif fmt in ('mp4', 'mov', 'mkv'):
        _run_ffmpeg(i + ['-c:v', 'libx264', '-preset', 'medium', '-crf', '20', '-pix_fmt', 'yuv420p',
                         '-vf', 'scale=trunc(iw/2)*2:trunc(ih/2)*2', '-c:a', 'aac', '-b:a', '192k']
                    + (['-movflags', '+faststart'] if fmt != 'mkv' else []) + [out])
    elif fmt == 'webm':
        _run_ffmpeg(i + ['-c:v', 'libvpx-vp9', '-crf', '32', '-b:v', '0', '-row-mt', '1', '-c:a', 'libopus', out])
    elif fmt == 'avi':
        _run_ffmpeg(i + ['-c:v', 'mpeg4', '-q:v', '4', '-c:a', 'libmp3lame', '-q:a', '3', out])
    elif fmt == 'mp3':
        _run_ffmpeg(i + ['-vn', '-c:a', 'libmp3lame', '-q:a', '2', out])
    elif fmt == 'wav':
        _run_ffmpeg(i + ['-vn', out])
    elif fmt == 'flac':
        _run_ffmpeg(i + ['-vn', '-c:a', 'flac', out])
    elif fmt == 'm4a':
        _run_ffmpeg(i + ['-vn', '-c:a', 'aac', '-b:a', '256k', out])
    elif fmt == 'ogg':
        _run_ffmpeg(i + ['-vn', '-c:a', 'libvorbis', '-q:a', '6', out])
    else:
        raise RuntimeError('不支持的目标格式:' + fmt)
    return out


def read_text(src):
    """把 txt / md / docx 读成 (文本, 是否Markdown)。"""
    ext = os.path.splitext(src)[1].lower()
    if ext == '.docx':
        import docx
        lines = []
        for p in docx.Document(src).paragraphs:
            st = (p.style.name or '').lower() if p.style is not None else ''
            if st.startswith('heading'):
                lv = ''.join(ch for ch in st if ch.isdigit()) or '1'
                lines.append('#' * min(6, int(lv)) + ' ' + p.text)
            elif st == 'title':
                lines.append('# ' + p.text)
            elif 'list' in st:
                lines.append('- ' + p.text)
            else:
                lines.append(p.text)
            lines.append('')
        return '\n'.join(lines), True
    raw = open(src, 'rb').read()
    for enc in ('utf-8-sig', 'utf-8', 'gb18030', 'utf-16'):
        try:
            return raw.decode(enc), ext in ('.md', '.markdown')
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', 'replace'), ext in ('.md', '.markdown')


def text_to_pdf(text, is_md, out):
    """用 Qt 的原生文本排版引擎排成 A4 PDF(需在界面线程调用)。"""
    from PySide6.QtCore import QMarginsF, QSizeF
    from PySide6.QtGui import QTextDocument, QPdfWriter, QPageSize, QPageLayout, QFont
    w = QPdfWriter(out)
    w.setPageSize(QPageSize(QPageSize.A4))
    w.setPageMargins(QMarginsF(18, 18, 18, 18), QPageLayout.Millimeter)
    w.setResolution(144)
    doc = QTextDocument()
    f = QFont()
    f.setPointSize(11)
    doc.setDefaultFont(f)
    if is_md:
        doc.setMarkdown(text)
    else:
        doc.setPlainText(text)
    doc.setPageSize(QSizeF(w.width(), w.height()))
    doc.print_(w)
    return out


def convert_doc(src, fmt):
    ext = os.path.splitext(src)[1].lower()
    out = _out(src, fmt)
    if ext == '.pdf':
        if fmt == 'txt':
            from pypdf import PdfReader
            with open(out, 'w', encoding='utf-8') as fh:
                for i, pg in enumerate(PdfReader(src).pages):
                    fh.write((pg.extract_text() or '') + '\n\n')
            return out
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(src)
        outs = []
        base = os.path.splitext(src)[0]
        for i in range(len(pdf)):
            im = pdf[i].render(scale=2).to_pil()
            o = unique_path(f"{base}_第{i + 1}页.{fmt}" if len(pdf) > 1 else f"{base}.{fmt}")
            (im.convert('RGB') if fmt == 'jpg' else im).save(o, quality=92) if fmt == 'jpg' else im.save(o)
            outs.append(o)
        return outs[0]
    text, is_md = read_text(src)
    if fmt == 'pdf':
        return text_to_pdf(text, is_md, out)
    if fmt in ('txt', 'md'):
        with open(out, 'w', encoding='utf-8') as fh:
            fh.write(text)
        return out
    if fmt == 'docx':
        import docx
        d = docx.Document()
        for line in text.splitlines():
            s = line.strip()
            if is_md and s.startswith('#'):
                lv = len(s) - len(s.lstrip('#'))
                d.add_heading(s.lstrip('#').strip(), level=min(lv, 6))
            elif is_md and (s.startswith('- ') or s.startswith('* ')):
                d.add_paragraph(s[2:], style='List Bullet')
            elif s:
                d.add_paragraph(line)
        d.save(out)
        return out
    raise RuntimeError('不支持的目标格式:' + fmt)


def needs_gui_thread(src, fmt):
    return fmt == 'pdf' and os.path.splitext(src)[1].lower() in ('.txt', '.md', '.markdown', '.docx')


def convert(src, fmt):
    cat = category(src)
    if cat == 'image':
        return convert_image(src, fmt)
    if cat in ('video', 'audio'):
        return convert_media(src, fmt)
    if cat == 'doc':
        return convert_doc(src, fmt)
    raise RuntimeError('三虎还不会转换这种文件')
